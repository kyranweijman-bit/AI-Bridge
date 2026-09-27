"""
Step 11: a Claude-style interface - a list of separate chats on the left,
file upload, and download buttons on every answer.

Install requirements first:  pip install -r requirements.txt
Run it with:                 streamlit run app.py

Streamlit reruns this whole file top-to-bottom on every click. That's why
state that needs to survive a click (which chat is open, its messages)
either lives in the database (see db.py) or in st.session_state, which
Streamlit itself preserves across reruns.
"""

import os
import streamlit as st
from dotenv import load_dotenv

MAX_TOTAL_FILE_CHARS = 60000  # combined cap across every file attached to one question, to control cost

# Load .env right away, before the password check below runs - bridge.py
# also calls this, but that import happens *after* the password gate, which
# meant APP_PASSWORD wasn't in os.environ yet when it was needed. Calling it
# here too fixes that (load_dotenv() only fills in variables that aren't
# already set, so calling it twice is harmless).
load_dotenv()


# ---- Step 12: pull secrets from Streamlit Cloud's secrets store, if present ----
#
# Locally, bridge.py's load_dotenv() reads .env instead - .env is never
# uploaded to GitHub, so it isn't available once hosted. On Streamlit Cloud,
# there is no .env file; instead you paste the same values into the app's
# "Secrets" box in its dashboard, and Streamlit exposes them as st.secrets.
# This copies them into the normal environment variables so bridge.py
# doesn't need to know or care which situation it's running in.
try:
    for key in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "CLAUDE_MODEL", "OPENAI_MODEL", "APP_PASSWORD",
                "DATABASE_URL", "SUPABASE_URL", "SUPABASE_SERVICE_KEY"):
        if key in st.secrets:
            # .strip() guards against a stray newline or trailing space that
            # can sneak in when copy-pasting a long key - those are invisible
            # to look at but break the actual API request.
            os.environ.setdefault(key, str(st.secrets[key]).strip())
except Exception:
    pass  # no secrets configured (e.g. running locally with just a .env file) - that's fine


# ---- Step 12: a simple password gate ----
#
# Without this, anyone who gets hold of the URL could ask questions on your
# API keys' bill. st.session_state.authenticated is what remembers "this
# browser already typed the right password" across reruns, for the rest of
# this browser tab's session.
def check_password():
    def password_entered():
        st.session_state.authenticated = (
            st.session_state.get("password_input") == os.environ.get("APP_PASSWORD")
        )

    if st.session_state.get("authenticated"):
        return True

    st.text_input("Password", type="password", on_change=password_entered, key="password_input")
    if st.session_state.get("authenticated") is False:
        st.error("Wrong password.")
    return False


st.set_page_config(page_title="AI Bridge", layout="wide")

if not check_password():
    st.stop()  # halts the script here - nothing below this line ever runs without the password

import bridge
import db
from bridge import (
    ask_claude,
    ask_claude_stream,
    ask_chatgpt_stream,
    ROLE_INSTRUCTIONS,
    REVIEW_INSTRUCTION,
    build_compare_prompt,
    build_debate_message,
    build_conclusion_prompt,
    FOLLOWUP_PRESETS,
    OTHER_MODEL_INSTRUCTION,
    PROMPT_PRESETS,
    RECIPE_STEPS,
    DEPTH_PRESETS,
    CONFIDENCE_TAG_INSTRUCTION,
    build_disagreement_map_prompt,
    build_challenge_assumptions_prompt,
    build_clarify_prompt,
    NO_CLARIFICATION_NEEDED,
    build_stall_check_prompt,
    build_consensus_check_prompt,
    STEELMAN_INSTRUCTION,
    TESTS_NOT_OPINIONS_INSTRUCTION,
    build_third_option_prompt,
    build_file_edit_instruction,
    extract_edited_files,
)
from pypdf import PdfReader
from docx import Document
from fpdf import FPDF
import html as html_lib
import random
import base64
import io


# ---- Step 24: multi-file upload with passage references ----
#
# Each file is broken into tagged chunks - by page for PDFs, by a group of
# paragraphs for .docx (which has no fixed notion of a "page"), by a group
# of lines for anything else - so the model can point back at *where* in
# the file something came from, e.g. "[report.pdf, page 3]", instead of
# treating the whole document as one undifferentiated blob of text.

def extract_tagged_text(uploaded_file):
    name = uploaded_file.name

    if name.lower().endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        chunks = []
        for i, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            if text:
                chunks.append(f"[{name}, page {i + 1}]\n{text}")
        return "\n\n".join(chunks)

    if name.lower().endswith(".docx"):
        document = Document(uploaded_file)
        paragraphs = [p.text for p in document.paragraphs]
        group_size = 20
        chunks = []
        for start in range(0, len(paragraphs), group_size):
            group = paragraphs[start:start + group_size]
            text = "\n".join(group).strip()
            if text:
                chunks.append(f"[{name}, paragraphs {start + 1}-{start + len(group)}]\n{text}")
        return "\n\n".join(chunks)

    text = uploaded_file.read().decode("utf-8", errors="replace")
    lines = text.splitlines()
    if not lines:
        return text
    group_size = 40
    chunks = []
    for start in range(0, len(lines), group_size):
        group = lines[start:start + group_size]
        chunk_text = "\n".join(group).strip()
        if chunk_text:
            chunks.append(f"[{name}, lines {start + 1}-{start + len(group)}]\n{chunk_text}")
    return "\n\n".join(chunks) if chunks else text



# ---- Image input ----
#
# An attached image isn't extracted as text like the other file types - it's
# base64-encoded into this app's neutral {"type": "image", ...} content block
# (see bridge.py's _to_claude_messages/_to_openai_messages) so both models
# can actually look at it. Capped at 5 MB each - big enough for a screenshot
# or phone photo, small enough to not balloon the request.
IMAGE_EXTENSIONS = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
                    ".webp": "image/webp", ".gif": "image/gif"}
MAX_IMAGE_BYTES = 5_000_000

# "Edit my file(s)" - a cap on how many attached files can be edited-and-
# downloaded at once, so the model isn't asked to juggle rewriting a dozen
# full files in one answer.
MAX_FILE_EDIT_FILES = 5


def build_files_prompt(uploaded_files, chat_id, turn_index, keep_originals=False):
    """Extracts, tags and caps every attached file's text, saves each one to
    the database, and returns (text_block, image_blocks) - text_block is the
    combined block to prepend to the prompt, image_blocks is a list of
    neutral image content blocks (empty if no images were attached). Any
    st.warning() calls happen here, right where the cap is applied.

    Image bytes are saved to the database like any other file, but the full
    base64 data is never folded into text_block - only a short placeholder
    line is, since text_block is what gets stored as this turn's permanent
    context and resent on every future follow-up question in the chat.
    Resending the full image on every later message would be a silent,
    ongoing token cost the user never asked for.

    keep_originals=True (set when "edit my file" is checked) also keeps each
    non-image file's original bytes in the database, so the edited download
    built from it later can carry over a .docx's styling or a .pdf's page
    size instead of starting from a blank page."""
    tagged_parts, total_len, truncated = [], 0, False
    image_blocks = []

    text_files = [uf for uf in uploaded_files if os.path.splitext(uf.name)[1].lower() not in IMAGE_EXTENSIONS]
    image_files = [uf for uf in uploaded_files if os.path.splitext(uf.name)[1].lower() in IMAGE_EXTENSIONS]

    for uf in text_files:
        text = extract_tagged_text(uf)
        remaining = MAX_TOTAL_FILE_CHARS - total_len
        if remaining <= 0:
            truncated = True
            break
        if len(text) > remaining:
            text = text[:remaining]
            truncated = True
        tagged_parts.append(text)
        total_len += len(text)

        warning = db.save_file(chat_id, turn_index, uf.name, text, uf.getvalue(), keep_original=keep_originals)
        if warning:
            st.warning(warning)

    for uf in image_files:
        data = uf.getvalue()
        if len(data) > MAX_IMAGE_BYTES:
            st.warning(f"{uf.name} is larger than 5 MB and was skipped - resize it and try again.")
            continue
        media_type = IMAGE_EXTENSIONS[os.path.splitext(uf.name)[1].lower()]
        image_blocks.append({
            "type": "image",
            "media_type": media_type,
            "data": base64.b64encode(data).decode("ascii"),
        })
        tagged_parts.append(f"[{uf.name} - image attached; not resent in later follow-up questions]")

        warning = db.save_file(chat_id, turn_index, uf.name, "(image attachment - see the file itself)", data)
        if warning:
            st.warning(warning)

    if truncated:
        st.warning(f"Attached files were longer than {MAX_TOTAL_FILE_CHARS:,} characters combined - only the start was sent.")

    names = ", ".join(uf.name for uf in uploaded_files)
    text_block = (
        f"The user attached {len(uploaded_files)} file(s): {names}.\n"
        "Each passage below is tagged with where it came from, like "
        "'[filename, page N]' - when you rely on specific content from "
        "these files, cite the tag it came from.\n\n"
        f"{chr(10).join(tagged_parts)}"
    )
    return text_block, image_blocks


# ---- Edit my file(s), and let me download them again ----
#
# Turns the model's answer text (already pulled out of the full answer by
# bridge.py's extract_edited_files()) into a downloadable file matching the
# original attachment's type. When the original file's bytes were kept
# (db.get_original_file_bytes - only true when "edit my file" was checked at
# upload time), a .docx reuses the original document as a template - each
# paragraph keeps its own style and its first run's formatting, just with
# its text replaced - and a .pdf reuses the original page size. This is a
# best-effort carry-over, not a byte-level edit: tables, headers/footers and
# inline images in the original are left as they were (nothing there gets
# text replaced), and if the edited text has a different number of
# paragraphs than the original, the extra ones are added plainly at the end
# (or the original's extra paragraphs are simply cleared). A .pdf's *text*
# is always freshly laid out - true fixed-layout PDF editing (reflowing
# edited text back into exact original positions) isn't something this app
# attempts, only its page size is carried over.
def _apply_edited_text_to_docx(doc, text):
    new_lines = text.split("\n")
    paras = doc.paragraphs
    for idx, para in enumerate(paras):
        new_text = new_lines[idx] if idx < len(new_lines) else ""
        if para.runs:
            para.runs[0].text = new_text
            for extra_run in para.runs[1:]:
                extra_run.text = ""
        else:
            para.add_run(new_text)
    for line in new_lines[len(paras):]:
        doc.add_paragraph(line)


def build_edited_file_bytes(text, original_filename, original_bytes=None):
    ext = os.path.splitext(original_filename)[1].lower()
    stem = os.path.splitext(original_filename)[0] or "edited"

    if ext == ".docx":
        doc = None
        if original_bytes:
            try:
                doc = Document(io.BytesIO(original_bytes))
                _apply_edited_text_to_docx(doc, text)
            except Exception:
                doc = None  # a corrupt/unreadable original - fall back below rather than fail the download
        if doc is None:
            doc = Document()
            for line in text.split("\n"):
                doc.add_paragraph(line)
        buf = io.BytesIO()
        doc.save(buf)
        return (
            buf.getvalue(),
            f"{stem}_edited.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    if ext == ".pdf":
        page_format = "A4"
        if original_bytes:
            try:
                reader = PdfReader(io.BytesIO(original_bytes))
                mediabox = reader.pages[0].mediabox
                # PDF points -> mm, what FPDF's format= wants.
                page_format = (float(mediabox.width) * 0.352778, float(mediabox.height) * 0.352778)
            except Exception:
                page_format = "A4"
        pdf = FPDF(format=page_format)
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        for line in text.split("\n"):
            # The built-in core font is Latin-1 only - swap anything it can't
            # render for "?" rather than let a stray smart-quote or emoji
            # crash the download.
            safe_line = line.encode("latin-1", "replace").decode("latin-1")
            pdf.multi_cell(0, 6, safe_line)
        return bytes(pdf.output()), f"{stem}_edited.pdf", "application/pdf"

    # Plain text, Markdown, or code (.py, .csv, ...) - same extension, as text.
    safe_ext = ext if ext else ".txt"
    return text.encode("utf-8"), f"{stem}_edited{safe_ext}", "text/plain"


# ---- Step 21: storage lives in a database (db.py) ----

@st.cache_resource
def connect_database():
    """Runs once per app start (not on every click): opens the connection
    and creates any missing tables."""
    db.init_schema()
    return True


try:
    connect_database()
except Exception as e:
    st.error(f"Couldn't connect to the database: {e}")
    st.info("Check DATABASE_URL in .env (locally) or in the app's Secrets on Streamlit Cloud. See SETUP_DATABASE.md.")
    st.stop()


def _log_usage(provider, model, input_tokens, output_tokens, cost):
    # st.session_state here is the session of whoever pressed Ask, because
    # this is called from inside that session's own script run.
    db.log_usage(provider, model, input_tokens, output_tokens, cost,
                 chat_id=st.session_state.get("current_chat_id"))


bridge.usage_callback = _log_usage


# ---- Pick which chat is open ----

if "current_chat_id" not in st.session_state:
    existing = db.list_chats()
    st.session_state.current_chat_id = existing[0][0] if existing else db.create_chat()

# Step 26: the multi-step flow currently running (if any) - which steps
# already succeeded, and which one (if any) just failed and is waiting to
# be retried. Step 22/27/30 reuse the same pattern for one-off actions
# (a conclusion, a follow-up, "add to memory") that don't need retry
# machinery of their own since they're a single call, not several in a row.
if "in_progress" not in st.session_state:
    st.session_state.in_progress = None
if "pending_action" not in st.session_state:
    st.session_state.pending_action = None


# ---- Sidebar: the chat list ----

st.sidebar.title("Chats")

if st.sidebar.button("+ New chat", use_container_width=True):
    st.session_state.current_chat_id = db.create_chat()
    st.session_state.in_progress = None
    st.rerun()

search = st.sidebar.text_input("Search chats", placeholder="Search titles and messages")

st.sidebar.divider()

chat_list = db.search_chats(search) if search else db.list_chats()
for chat_id, title in chat_list:
    label = title or "New chat"
    if chat_id == st.session_state.current_chat_id:
        label = f"-> {label}"

    col1, col2 = st.sidebar.columns([5, 1])
    with col1:
        if st.button(label, key=f"open_{chat_id}", use_container_width=True):
            st.session_state.current_chat_id = chat_id
            st.session_state.in_progress = None
            st.rerun()
    with col2:
        if st.button("x", key=f"del_{chat_id}"):
            db.delete_chat(chat_id)
            if st.session_state.current_chat_id == chat_id:
                remaining = db.list_chats()
                st.session_state.current_chat_id = remaining[0][0] if remaining else db.create_chat()
            st.rerun()

if search and not chat_list:
    st.sidebar.caption("No chats match.")


# ---- Step 21: memory ----
#
# Memories are sent to both models as a standing instruction on every call.
# "All chats" ones apply everywhere; "This chat only" ones just here.
st.sidebar.divider()
with st.sidebar.expander("Memory"):
    for m in db.list_memories(st.session_state.current_chat_id):
        scope = "all chats" if m["chat_id"] is None else "this chat"
        mcol1, mcol2 = st.columns([5, 1])
        mcol1.caption(f"{m['content']}  \n*({scope})*")
        if mcol2.button("x", key=f"mem_del_{m['id']}"):
            db.delete_memory(m["id"])
            st.rerun()

    with st.form("add_memory", clear_on_submit=True):
        new_memory = st.text_area("Remember...", placeholder="e.g. I study at HAN; answer in English")
        memory_scope = st.radio("Applies to", ["All chats", "This chat only"], horizontal=True)
        if st.form_submit_button("Save to memory") and new_memory.strip():
            db.add_memory(
                new_memory,
                chat_id=None if memory_scope == "All chats" else st.session_state.current_chat_id,
            )
            st.rerun()


# ---- Editable working brief: a visible summary of the goal, constraints,
# decisions and open questions for this chat, sent to both models on every
# call (see the `memory` merge in the main area below), that can be
# corrected at any time - unlike memory, one editable block rather than a
# growing list of separate notes. ----
st.sidebar.divider()
with st.sidebar.expander("Working brief (this chat)"):
    _current_brief = db.get_brief(st.session_state.current_chat_id)
    _new_brief = st.text_area(
        "Goal, constraints, decisions, open questions...",
        value=_current_brief, height=150, key=f"brief_{st.session_state.current_chat_id}",
    )
    if st.button("Save brief", key="save_brief"):
        db.set_brief(st.session_state.current_chat_id, _new_brief)
        st.rerun()
    st.caption("Sent to both models on every call in this chat, like memory - but as one editable block.")


# ---- More roles and personas: saveable personas (a name + a free-text
# instruction) assignable to either model's Debate role, alongside the fixed
# Proposer / Critic / Fact-checker / Devil's advocate roles. ----
st.sidebar.divider()
with st.sidebar.expander("Personas"):
    for p in db.list_personas():
        ppcol1, ppcol2 = st.columns([5, 1])
        ppcol1.caption(f"**{p['name']}** - {p['instruction']}")
        if ppcol2.button("x", key=f"persona_del_{p['id']}"):
            db.delete_persona(p["id"])
            st.rerun()

    with st.form("add_persona", clear_on_submit=True):
        new_persona_name = st.text_input("Name", placeholder="e.g. Strict teacher")
        new_persona_instruction = st.text_area("Instruction", placeholder="e.g. Grade harshly and demand rigor; never soften feedback.")
        if st.form_submit_button("Save persona") and new_persona_name.strip() and new_persona_instruction.strip():
            db.add_persona(new_persona_name, new_persona_instruction)
            st.rerun()
    st.caption("Saved personas appear as extra role options in Debate mode.")


# ---- Decision journal: what you chose, why, and (filled in later) what
# actually happened - logged per turn further down, browsed here across
# every chat. ----
st.sidebar.divider()
with st.sidebar.expander("Decision journal"):
    _all_decisions = db.list_decisions()
    if not _all_decisions:
        st.caption('Nothing logged yet - use "Log a decision from this turn" under any answer.')
    for d in _all_decisions:
        st.caption(f"**{d['title']}** - {d['choice']}" + (f" *({d['reasoning']})*" if d["reasoning"] else ""))
        _new_outcome = st.text_input("Outcome (fill in later)", value=d.get("outcome") or "", key=f"journal_outcome_{d['id']}")
        if st.button("Save outcome", key=f"journal_save_{d['id']}"):
            db.set_decision_outcome(d["id"], _new_outcome.strip())
            st.rerun()
        st.divider()


# ---- Model-vs-model stats dashboard, from blind-judging votes. ----
st.sidebar.divider()
with st.sidebar.expander("Model stats"):
    _vote_stats = db.vote_stats()
    if not _vote_stats:
        st.caption('No blind-judging votes yet - turn on "Blind judging" on an Independent Answers question to start collecting them.')
    else:
        _total_votes = sum(_vote_stats.values())
        for k, v in _vote_stats.items():
            st.write(f"{k}: {v} ({v / _total_votes * 100:.0f}%)")
        st.caption("Overall only, across every chat - a per-topic breakdown isn't built yet.")


# ---- Step 16/21: cost visibility, now saved permanently ----
st.sidebar.divider()
with st.sidebar.expander("Usage"):
    usage = db.usage_summary(st.session_state.current_chat_id)
    st.write(f"Today: ≈ ${usage['today_cost']:.4f}")
    st.write(f"This chat: ≈ ${usage['chat_cost']:.4f}")
    st.markdown(f"**All time: ≈ ${usage['total_cost']:.4f}**")
    st.caption(f"{usage['total_tokens']:,} tokens in total")
    for row in db.usage_by_model():
        cost = f"≈ ${float(row['cost']):.4f}" if row["cost"] is not None else "no price known"
        st.caption(f"{row['model']}: {int(row['input_tokens']):,} in / {int(row['output_tokens']):,} out, {cost}")
    if usage["unpriced_calls"]:
        st.caption("Some calls have no cost estimate - their model isn't in the PRICING table in bridge.py.")

    # Spending limit: stored in the settings table so it's shared and
    # survives restarts, not just this browser tab. Checked before every
    # model call (see spending_limit_status() and its use in run_step()
    # further down) - once today's cost reaches it, further calls are
    # blocked (with a clear reason) until it's raised or the day rolls over.
    st.divider()
    current_limit = db.get_setting("daily_spending_limit", "")
    new_limit = st.text_input("Daily spending limit ($, blank = no limit)", value=current_limit, key="spending_limit_input")
    if st.button("Save limit"):
        db.set_setting("daily_spending_limit", new_limit.strip())
        st.rerun()


def spending_limit_status():
    """None if there's no limit or it hasn't been reached; otherwise
    (today_cost, limit) - both floats - so the caller can show why."""
    raw = db.get_setting("daily_spending_limit", "")
    if not raw:
        return None
    try:
        limit = float(raw)
    except ValueError:
        return None  # something non-numeric got saved - treat as "no limit" rather than crash
    today_cost = db.usage_summary(st.session_state.current_chat_id)["today_cost"]
    if today_cost >= limit:
        return today_cost, limit
    return None


# ---- Main area: the open chat ----

chat = db.load_chat(st.session_state.current_chat_id)
if chat is None:
    # The chat this browser tab remembered was deleted (e.g. in another
    # tab) - start a fresh one instead of crashing.
    st.session_state.current_chat_id = db.create_chat()
    chat = db.load_chat(st.session_state.current_chat_id)

chat_id = st.session_state.current_chat_id
turn_index = len(chat["display"])
memory = db.memory_prompt(chat_id)  # sent to both models as a system prompt

# Editable working brief: folded into the same system-prompt slot as memory,
# just as one editable block instead of a growing list of separate notes.
_brief = db.get_brief(chat_id)
if _brief:
    _brief_block = f"Working brief for this chat:\n{_brief}"
    memory = f"{memory}\n\n{_brief_block}" if memory else _brief_block

chat_decisions = db.list_decisions(chat_id)  # loaded once per rerun, filtered per turn below


def turn_transcript_text(turn):
    """Step 22/30: the whole "discussion" a turn represents, as plain text -
    used both to ask for a final conclusion and to fill in a quick
    "add to memory" save with something more useful than a raw dump."""
    t = turn.get("type", "single")
    if t == "single":
        primary = turn.get("primary", "Claude")
        reviewer = "ChatGPT" if primary == "Claude" else "Claude"
        first = turn["claude"] if primary == "Claude" else turn["chatgpt"]
        second = turn["chatgpt"] if primary == "Claude" else turn["claude"]
        return f"{primary} answered:\n{first}\n\n{reviewer} reviewed:\n{second}"
    if t == "independent":
        return (
            f"Claude answered:\n{turn['claude']}\n\n"
            f"ChatGPT answered:\n{turn['chatgpt']}\n\n"
            f"Comparison:\n{turn['comparison']}"
        )
    if t in ("debate", "recipe"):
        parts = []
        for entry in turn["transcript"]:
            if len(entry) == 3:
                speaker, label, reply = entry
                parts.append(f"{speaker} ({label}): {reply}")
            else:
                speaker, reply = entry
                parts.append(f"{speaker}: {reply}")
        return "\n\n".join(parts)
    return turn.get("answer", "")  # solo


# ---- Quick win: confidence-tag highlighting ----
#
# When a turn was asked with confidence tags on, each model's answer has
# inline "[sure]" / "[fairly sure]" / "[guessing]" markers in it (see
# CONFIDENCE_TAG_INSTRUCTION in bridge.py). This turns those into small
# colored badges instead of leaving them as plain bracketed text. The rest
# of the answer is HTML-escaped first, so this only ever adds the fixed,
# safe <span> markup below - nothing from the model's own text can inject
# further HTML.
CONFIDENCE_TAG_STYLES = {
    "[sure]": ("#1a7f37", "sure"),
    "[fairly sure]": ("#9a6700", "fairly sure"),
    "[guessing]": ("#cf222e", "guessing"),
}


def render_with_confidence_tags(text):
    escaped = html_lib.escape(text).replace("\n", "<br>")
    for tag, (color, label) in CONFIDENCE_TAG_STYLES.items():
        badge = (
            f'<span style="background:{color}22;color:{color};border-radius:4px;'
            f'padding:1px 6px;font-size:0.78em;font-weight:600;">{label}</span>'
        )
        escaped = escaped.replace(html_lib.escape(tag), badge)
    st.markdown(escaped, unsafe_allow_html=True)


# ---- Quick win: copy as Markdown / LaTeX ----
#
# st.code() blocks get a copy-to-clipboard icon for free in Streamlit, so
# "Markdown" view is just the raw text in a code block; "LaTeX" wraps it in
# a minimal document so it's ready to paste into an existing .tex file.
def show_text(text, view_mode, tagged=False):
    if view_mode == "Markdown":
        st.code(text, language="markdown")
    elif view_mode == "LaTeX":
        st.code(f"\\documentclass{{article}}\n\\begin{{document}}\n{text}\n\\end{{document}}", language="latex")
    elif tagged:
        render_with_confidence_tags(text)
    else:
        st.write(text)


title_col, pin_col = st.columns([8, 1])
title_col.title(chat["title"])
with pin_col:
    if st.button("Unpin" if chat["pinned"] else "Pin", help="Pinned chats stay at the top of the list"):
        db.set_pinned(chat_id, not chat["pinned"])
        st.rerun()

with st.expander("Rename chat"):
    new_title = st.text_input("New name", value=chat["title"], key=f"rename_{chat_id}")
    if st.button("Save name") and new_title.strip():
        db.rename_chat(chat_id, new_title.strip()[:100])
        st.rerun()

# Step 25: export the whole conversation - every question, its attachments,
# and every answer/review/comparison/debate-or-recipe reply/conclusion - as
# one readable Markdown file, instead of only individual answers.
st.download_button(
    "Export whole conversation",
    db.export_chat_text(chat_id),
    file_name=f"{(chat['title'] or 'chat').strip()[:50]}.md",
    key="export_chat",
)


# =====================================================================
# Either: show the "ask a new question" form (mode picker, presets, etc.)
# and handle any one-off action a button lower down just queued up - or,
# if a multi-step flow is currently running/paused on an error, show its
# progress/retry controls instead. Never both at once, so a half-finished
# flow can't be interrupted by starting a second one.
# =====================================================================

if st.session_state.in_progress is None:

    uploaded_files = st.file_uploader(
        "Attach file(s) (optional) - text, PDF, Word, or an image (PNG/JPG/WEBP/GIF)",
        accept_multiple_files=True,
    )
    question = st.text_input("Ask a question")

    # Depth control: one toggle that sets mode + rounds + token budget
    # together, each preset bounded in both rounds and tokens so it can't
    # accidentally turn into a long, expensive run. A prompt preset (below)
    # already picks its own mode/settings, so it takes priority when both
    # are set - the caption makes it clear which one is actually driving.
    depth_name = st.selectbox("Depth (optional)", ["Custom (manual controls)"] + list(DEPTH_PRESETS.keys()))
    depth = DEPTH_PRESETS.get(depth_name)

    # Step 29: one-click templates. Picking one fixes the mode (and roles,
    # for Debate) to whatever suits that kind of question, and prepends a
    # short instruction to whatever you actually typed - your own question
    # is always kept, never replaced. Pick "None" for full manual control.
    preset_name = st.selectbox("Quick start (optional)", ["None"] + list(PROMPT_PRESETS.keys()))
    preset = PROMPT_PRESETS.get(preset_name)

    MODE_OPTIONS = ["Single review", "Independent answers", "Debate",
                    "Recipe: Draft -> Critique -> Revise -> Check", "One model only"]

    primary = rounds = claude_role = chatgpt_role = which_model = debate_start = None
    max_tokens = None
    pause_between_rounds = False
    ask_before_debating = False
    debate_length_mode = "Fixed rounds"
    steelman = False

    if preset:
        mode = preset["mode"]
        primary = preset.get("primary", "Claude")
        rounds = preset.get("rounds", 3)
        claude_role = preset.get("claude_role", "Proposer")
        chatgpt_role = preset.get("chatgpt_role", "Critic")
        which_model = preset.get("which_model", "Claude")
        debate_start = "Claude"
        detail = f" ({primary} first)" if mode == "Single review" else (f" ({which_model})" if mode == "One model only" else "")
        st.caption(f"Preset **{preset_name}** -> {mode}{detail}.")

    elif depth:
        mode = depth["mode"]
        max_tokens = depth["max_tokens"]
        primary = depth.get("primary", "Claude")
        rounds = depth.get("rounds", 3)
        claude_role = depth.get("claude_role", "Proposer")
        chatgpt_role = depth.get("chatgpt_role", "Critic")
        which_model = depth.get("which_model", "Claude")
        debate_start = "Claude"
        st.caption(f"Depth **{depth_name}** -> {mode}, up to {max_tokens:,} tokens per reply.")

    else:
        mode = st.radio("Mode", MODE_OPTIONS, horizontal=True)
        if mode == "Single review":
            primary = st.radio("Who answers first?", ["Claude", "ChatGPT"], horizontal=True)
        elif mode == "Debate":
            rounds = st.slider("Rounds", min_value=1, max_value=6, value=3)
            debate_start = st.radio("Who starts?", ["Claude", "ChatGPT"], horizontal=True)
            role_options = list(ROLE_INSTRUCTIONS.keys()) + [p["name"] for p in db.list_personas()]
            role_col1, role_col2 = st.columns(2)
            with role_col1:
                claude_role = st.selectbox("Claude's role", role_options, index=0)
            with role_col2:
                chatgpt_role = st.selectbox("ChatGPT's role", role_options, index=1)
            pause_between_rounds = st.checkbox(
                "Pause after each round so I can stop or steer it", value=True,
                help="Uncheck to run straight through all rounds automatically, like before.",
            )
            ask_before_debating = st.checkbox(
                "Ask a clarifying question first, if one would matter",
                help="One quick check before round 1: if an important detail is missing, you're asked "
                     "for it before the debate starts; if not, it carries straight on.",
            )
            debate_length_mode = st.radio(
                "Debate length", [
                    "Fixed rounds",
                    "Stop early once replies stop adding new info",
                    "Consensus-or-bust (keep going until both sides agree)",
                ],
                help="Either way, it never runs more than the rounds set above - this only ever stops it earlier.",
            )
            steelman = st.checkbox(
                "Steelman first - each side restates the other's point fairly before responding",
            )
        elif mode == "One model only":
            which_model = st.radio("Which model?", ["Claude", "ChatGPT"], horizontal=True)

    confidence_tags = st.checkbox(
        "Add confidence tags (sure / fairly sure / guessing) to claims",
        help="Each model tags its own claims inline; tags are shown as small colored badges once the answer is saved.",
    )
    challenge_assumptions = st.checkbox(
        "Challenge my assumptions first",
        help="Before answering, Claude points out questionable assumptions in your question - including ones you might share without noticing.",
    )
    tests_not_opinions = st.checkbox(
        "Prefer tests over opinions where the question is checkable",
        help="For checkable questions, the models are asked to propose a calculation, experiment or small test that would settle it, rather than just asserting an opinion.",
    )

    # "Edit my file(s), and let me download them again": offered for up to
    # MAX_FILE_EDIT_FILES non-image attachments at once (several is fine now
    # - each gets its own marker in the answer so they can be told apart).
    # Images are out of scope here (that's a different kind of editing).
    non_image_files = [uf for uf in (uploaded_files or []) if os.path.splitext(uf.name)[1].lower() not in IMAGE_EXTENSIONS]
    file_edit_mode = False
    if 1 <= len(non_image_files) <= MAX_FILE_EDIT_FILES:
        names_label = non_image_files[0].name if len(non_image_files) == 1 \
            else f"{len(non_image_files)} files: " + ", ".join(uf.name for uf in non_image_files)
        file_edit_mode = st.checkbox(
            f"Ask for the edited file(s) back ({names_label})",
            help='The model is asked to return each attached file\'s complete corrected content; '
                 'once it answers, a "Download edited file" button appears below for each one - '
                 "same format when possible (.docx keeps the original's styles, .pdf keeps its "
                 "page size, plain text/code keeps its extension) rather than a byte-level edit of "
                 "the original. Works in every mode: One model only and Single review edit "
                 "directly (Single review uses the reviewer's polished answer); Independent "
                 "answers gives you back Claude's AND ChatGPT's edited versions separately, so you "
                 "can pick; Debate and Recipe run one extra quick call at the end to produce the "
                 "final edited file(s) from the whole discussion, instead of repeating the full "
                 "file in every round.",
        )
    elif len(non_image_files) > MAX_FILE_EDIT_FILES:
        st.caption(f'Attach at most {MAX_FILE_EDIT_FILES} non-image files to also get a "Download edited file" option for each.')

    blind_judging = False
    if mode == "Independent answers":
        blind_judging = st.checkbox(
            "Blind judging - hide which model wrote which until I vote",
            help="Shows the two answers as \"Answer A\" / \"Answer B\"; your vote is recorded before the real labels are revealed.",
        )

    if st.button("Ask") and question:
        limit_hit = spending_limit_status()
        if limit_hit is not None:
            st.error(f"Daily spending limit reached (${limit_hit[0]:.2f} of ${limit_hit[1]:.2f}). Raise it in the sidebar's Usage panel, or wait until tomorrow.")
        else:
            prompt = question
            image_blocks = []

            if uploaded_files:
                files_block, image_blocks = build_files_prompt(
                    uploaded_files, chat_id, turn_index, keep_originals=file_edit_mode,
                )
                prompt = f"{files_block}\n\nQuestion: {question}"

            if preset:
                prompt = f"{preset['prefix']}\n\n{prompt}"

            # The plain-text version is always what's stored as this turn's
            # context (see db.add_turn below) - a follow-up question resends
            # this, never the raw image bytes.
            prompt_for_storage = prompt

            if image_blocks and mode in ("Single review", "Independent answers", "One model only"):
                prompt = [{"type": "text", "text": prompt}] + image_blocks
            elif image_blocks:
                st.info(
                    "Images are only sent as visual input in Single review, Independent answers, "
                    "and One model only mode for now - they were still saved as attachments to this chat."
                )

            st.session_state.in_progress = {
                "mode": mode, "question": question, "prompt": prompt,
                "prompt_for_storage": prompt_for_storage, "turn_index": turn_index,
                "primary": primary, "rounds": rounds, "claude_role": claude_role,
                "chatgpt_role": chatgpt_role, "which_model": which_model, "debate_start": debate_start,
                "max_tokens": max_tokens, "confidence_tags": confidence_tags,
                "blind_judging": blind_judging, "pause_between_rounds": pause_between_rounds,
                "challenge_assumptions": challenge_assumptions, "tests_not_opinions": tests_not_opinions,
                "ask_before_debating": ask_before_debating, "debate_length_mode": debate_length_mode,
                "steelman": steelman,
                "file_edit_mode": file_edit_mode,
                "file_edit_filenames": [uf.name for uf in non_image_files] if file_edit_mode else [],
                "steps": {}, "error": None, "paused": False, "next_round_note": "",
                "clarification_done": False,
            }
            st.rerun()

    # ---- Step 22/27/30: run whichever one-off action a button queued up ----
    pending = st.session_state.pending_action
    if pending is not None:
        st.session_state.pending_action = None
        target = chat["display"][pending["turn_index"]] if pending["turn_index"] < len(chat["display"]) else None

        if target is None:
            pass  # the turn it referred to is gone (e.g. chat changed) - just drop it

        elif pending["kind"] == "conclusion" and spending_limit_status() is not None:
            hit = spending_limit_status()
            st.error(f"Daily spending limit reached (${hit[0]:.2f} of ${hit[1]:.2f}). Raise it in the sidebar's Usage panel, or wait until tomorrow.")

        elif pending["kind"] == "conclusion":
            discussion = turn_transcript_text(target)
            conclusion_prompt = build_conclusion_prompt(target["question"], discussion)
            try:
                st.markdown("**Conclusion:**")
                conclusion_text = st.write_stream(
                    ask_claude_stream([{"role": "user", "content": conclusion_prompt}], system=memory)
                )
                db.add_extra(chat_id, pending["turn_index"], target.get("type", "single"),
                             "conclusion", "Claude", conclusion_text, in_context=True)
            except Exception as e:
                st.error(f"Couldn't generate a conclusion: {e}")
            st.rerun()

        elif pending["kind"] == "memory":
            text = (target.get("conclusion") or turn_transcript_text(target))[:2000]
            db.add_memory(text, chat_id=chat_id)
            st.rerun()

        elif pending["kind"] == "disagreement_map" and spending_limit_status() is not None:
            hit = spending_limit_status()
            st.error(f"Daily spending limit reached (${hit[0]:.2f} of ${hit[1]:.2f}). Raise it in the sidebar's Usage panel, or wait until tomorrow.")

        elif pending["kind"] == "disagreement_map":
            discussion = turn_transcript_text(target)
            dmap_prompt = build_disagreement_map_prompt(target["question"], discussion)
            try:
                st.markdown("**Disagreement map:**")
                dmap_text = st.write_stream(
                    ask_claude_stream([{"role": "user", "content": dmap_prompt}], system=memory)
                )
                db.add_extra(chat_id, pending["turn_index"], target.get("type", "single"),
                             "disagreement_map", "Claude", dmap_text, in_context=False)
            except Exception as e:
                st.error(f"Couldn't generate a disagreement map: {e}")
            st.rerun()

        elif pending["kind"] == "third_option" and spending_limit_status() is not None:
            hit = spending_limit_status()
            st.error(f"Daily spending limit reached (${hit[0]:.2f} of ${hit[1]:.2f}). Raise it in the sidebar's Usage panel, or wait until tomorrow.")

        elif pending["kind"] == "third_option":
            discussion = turn_transcript_text(target)
            third_prompt = build_third_option_prompt(target["question"], discussion)
            try:
                st.markdown("**A third option:**")
                third_text = st.write_stream(
                    ask_claude_stream([{"role": "user", "content": third_prompt}], system=memory)
                )
                db.add_extra(chat_id, pending["turn_index"], target.get("type", "single"),
                             "third_option", "Claude", third_text, in_context=False)
            except Exception as e:
                st.error(f"Couldn't come up with a third option: {e}")
            st.rerun()

        elif pending["kind"] == "followup" and spending_limit_status() is not None:
            hit = spending_limit_status()
            st.error(f"Daily spending limit reached (${hit[0]:.2f} of ${hit[1]:.2f}). Raise it in the sidebar's Usage panel, or wait until tomorrow.")

        elif pending["kind"] == "followup":
            t = target.get("type", "single")
            if t == "solo":
                acting_model = target["model"]
            elif t == "single":
                acting_model = target.get("primary", "Claude")
            else:  # independent - the comparison (last thing in shared history) is always Claude's
                acting_model = "Claude"

            action = pending["action"]
            if action == "other_model":
                speak_model = "ChatGPT" if acting_model == "Claude" else "Claude"
                instruction = OTHER_MODEL_INSTRUCTION
            else:
                speak_model = acting_model
                instruction = FOLLOWUP_PRESETS[action]

            followup_messages = chat["history"] + [{"role": "user", "content": instruction}]
            try:
                st.markdown(f"**{speak_model}:**")
                stream = ask_claude_stream(followup_messages, system=memory) if speak_model == "Claude" \
                    else ask_chatgpt_stream(followup_messages, system=memory)
                answer = st.write_stream(stream)
                new_turn = {"type": "solo", "question": instruction, "model": speak_model, "answer": answer}
                db.add_turn(chat_id, turn_index, new_turn, instruction)
            except Exception as e:
                st.error(f"That follow-up failed: {e}")
            st.rerun()

else:
    # ---- Step 26: a multi-step flow is running, or paused on a failed step ----
    progress = st.session_state.in_progress

    if progress.get("error"):
        st.error(f"That step failed: {progress['error']}")
        retry_col, cancel_col = st.columns(2)
        if retry_col.button("Retry"):
            progress["error"] = None
            st.session_state.in_progress = progress
            st.rerun()
        if cancel_col.button("Cancel this question"):
            st.session_state.in_progress = None
            st.rerun()
        st.stop()  # don't attempt anything else until Retry is pressed

    def run_step(step_key, label, make_stream):
        """Runs one model call within the flow below. A step that already
        succeeded (on an earlier attempt, before something else failed)
        replays instantly from cache instead of calling the model again -
        only the step that actually failed gets retried. A step that raises,
        or the spending limit being hit, stops the whole script run here;
        the error banner above picks it up on the next run."""
        st.markdown(label)
        if step_key in progress["steps"]:
            st.write(progress["steps"][step_key])
            return progress["steps"][step_key]

        limit_hit = spending_limit_status()
        if limit_hit is not None:
            progress["error"] = (
                f"Daily spending limit reached (${limit_hit[0]:.2f} of ${limit_hit[1]:.2f}) "
                "- raise it in the sidebar's Usage panel, or wait until tomorrow."
            )
            st.session_state.in_progress = progress
            st.rerun()

        try:
            result = st.write_stream(make_stream())
        except Exception as e:
            progress["error"] = str(e)
            st.session_state.in_progress = progress
            st.rerun()
        progress["steps"][step_key] = result
        st.session_state.in_progress = progress
        return result

    def run_quiet_step(step_key, call_fn):
        """Like run_step, but for a small non-streamed check call (e.g. a
        yes/no "have they stalled?" check) that shouldn't clutter the screen
        with its own streamed output. Same caching/spending-limit/retry
        behavior, just quiet - and it calls a plain ask_claude(), not a
        *_stream() generator."""
        if step_key in progress["steps"]:
            return progress["steps"][step_key]

        limit_hit = spending_limit_status()
        if limit_hit is not None:
            progress["error"] = (
                f"Daily spending limit reached (${limit_hit[0]:.2f} of ${limit_hit[1]:.2f}) "
                "- raise it in the sidebar's Usage panel, or wait until tomorrow."
            )
            st.session_state.in_progress = progress
            st.rerun()

        try:
            result = call_fn()
        except Exception as e:
            progress["error"] = str(e)
            st.session_state.in_progress = progress
            st.rerun()
        progress["steps"][step_key] = result
        st.session_state.in_progress = progress
        return result

    mode = progress["mode"]
    question = progress["question"]
    prompt = progress["prompt"]
    flow_turn_index = progress["turn_index"]
    mt = progress.get("max_tokens")  # None = bridge.py's own default (MAX_TOKENS)
    sys_prompt = memory
    if progress.get("confidence_tags"):
        sys_prompt = f"{memory}\n\n{CONFIDENCE_TAG_INSTRUCTION}" if memory else CONFIDENCE_TAG_INSTRUCTION
    if progress.get("tests_not_opinions"):
        sys_prompt = f"{sys_prompt}\n\n{TESTS_NOT_OPINIONS_INSTRUCTION}" if sys_prompt else TESTS_NOT_OPINIONS_INSTRUCTION
    # "Edit my file(s)": for these three modes, every call already produces
    # one full candidate answer anyway, so folding the instruction into every
    # call is free. Debate and Recipe are handled differently, further down
    # below (one dedicated wrap-up call after the discussion) rather than
    # asking every round/step to re-embed the whole file(s) every time.
    if progress.get("file_edit_mode") and mode in ("Single review", "Independent answers", "One model only"):
        file_edit_instr = build_file_edit_instruction(progress["file_edit_filenames"])
        sys_prompt = f"{sys_prompt}\n\n{file_edit_instr}" if sys_prompt else file_edit_instr
    new_turn = None

    # "Challenge my assumptions": a quick, informational pre-step ahead of
    # whichever mode runs below - doesn't change `prompt` itself, just shows
    # what it found before the main answer(s).
    if progress.get("challenge_assumptions"):
        run_step(
            "challenge", "**Checking your question for questionable assumptions:**",
            lambda: ask_claude_stream([{"role": "user", "content": build_challenge_assumptions_prompt(question)}], system=sys_prompt),
        )

    if mode == "Single review":
        primary = progress["primary"]
        history_ctx = chat["history"] + [{"role": "user", "content": prompt}]
        first_answer = run_step(
            "first", f"**{primary}** (answering):",
            lambda: (ask_claude_stream if primary == "Claude" else ask_chatgpt_stream)(history_ctx, system=sys_prompt, max_tokens=mt),
        )
        reviewer = "ChatGPT" if primary == "Claude" else "Claude"
        reviewer_messages = history_ctx + [
            {"role": "assistant", "content": first_answer},
            {"role": "user", "content": REVIEW_INSTRUCTION},
        ]
        second_answer = run_step(
            "second", f"**{reviewer}** (reviewing):",
            lambda: (ask_chatgpt_stream if reviewer == "ChatGPT" else ask_claude_stream)(reviewer_messages, system=sys_prompt, max_tokens=mt),
        )
        new_turn = {
            "type": "single", "question": question, "primary": primary,
            "claude": first_answer if primary == "Claude" else second_answer,
            "chatgpt": first_answer if primary == "ChatGPT" else second_answer,
        }

    elif mode == "Independent answers":
        history_ctx = chat["history"] + [{"role": "user", "content": prompt}]
        claude_answer = run_step("claude", "**Claude** (answering independently):",
                                  lambda: ask_claude_stream(history_ctx, system=sys_prompt, max_tokens=mt))
        chatgpt_answer = run_step("chatgpt", "**ChatGPT** (answering independently):",
                                   lambda: ask_chatgpt_stream(history_ctx, system=sys_prompt, max_tokens=mt))
        compare_prompt = build_compare_prompt(question, claude_answer, chatgpt_answer)
        comparison = run_step("comparison", "**Comparison:**",
                               lambda: ask_claude_stream([{"role": "user", "content": compare_prompt}], system=sys_prompt, max_tokens=mt))
        new_turn = {
            "type": "independent", "question": question,
            "claude": claude_answer, "chatgpt": chatgpt_answer, "comparison": comparison,
        }

    elif mode == "Debate":
        # Step: stop and intervene. With "pause after each round" on, only
        # one round runs per script pass - after it, the flow parks itself
        # (progress["paused"] = True) and shows Continue / Stop here /
        # Cancel instead of ploughing straight through every round. "Stop
        # here" just shrinks `rounds` down to what's already done, so the
        # normal finalize-the-turn path below picks it up unchanged.
        rounds = progress["rounds"]
        roles = {"claude": progress["claude_role"], "chatgpt": progress["chatgpt_role"]}
        pause_on = progress.get("pause_between_rounds", False)
        length_mode = progress.get("debate_length_mode", "Fixed rounds")
        # More roles and personas: saved personas (db.py's personas table)
        # are looked up alongside the fixed Proposer/Critic/Fact-checker/
        # Devil's advocate role instructions, so either can be assigned here.
        all_roles = {**ROLE_INSTRUCTIONS, **{p["name"]: p["instruction"] for p in db.list_personas()}}

        # "Ask before debating": one quiet check before round 1 only - once
        # it's been done (either answer given, skipped, or no clarification
        # was needed), it never runs again for this question.
        if progress.get("ask_before_debating") and not progress.get("clarification_done"):
            clarify_text = run_step(
                "clarify", "**Checking whether a clarifying question is needed first:**",
                lambda: ask_claude_stream([{"role": "user", "content": build_clarify_prompt(question)}], system=sys_prompt),
            )
            if NO_CLARIFICATION_NEEDED in clarify_text.lower():
                progress["clarification_done"] = True
                st.session_state.in_progress = progress
            else:
                st.info("Before debating, one clarifying question:")
                clarify_answer = st.text_area("Your answer (optional - leave blank to skip)", key="clarify_answer")
                cc1, cc2 = st.columns(2)
                if cc1.button("Continue with this answer"):
                    progress["clarification_done"] = True
                    if clarify_answer.strip():
                        note = f"\n\nClarification from the user: {clarify_answer.strip()}"
                        if isinstance(progress["prompt"], str):
                            progress["prompt"] += note
                        else:
                            progress["prompt"][0]["text"] += note
                    st.session_state.in_progress = progress
                    st.rerun()
                if cc2.button("Skip and continue"):
                    progress["clarification_done"] = True
                    st.session_state.in_progress = progress
                    st.rerun()
                st.stop()

        transcript = []
        speaker = "claude" if progress["debate_start"] == "Claude" else "chatgpt"
        for i in range(rounds):
            key = f"round_{i}"
            if key not in progress["steps"]:
                break
            speaker_name = "Claude" if speaker == "claude" else "ChatGPT"
            role = roles["claude"] if speaker_name == "Claude" else roles["chatgpt"]
            st.markdown(f"**{speaker_name}** (round {i + 1}, {role}):")
            st.write(progress["steps"][key])
            transcript.append((speaker_name, progress["steps"][key]))
            speaker = "chatgpt" if speaker == "claude" else "claude"
        completed = len(transcript)

        if pause_on and progress.get("paused") and completed < rounds:
            st.info(f"Paused after round {completed} of {rounds}. Continue, add a note for the next round, or stop here.")
            note = st.text_area("Add a constraint or correction before continuing (optional)", key="debate_note")
            c1, c2, c3 = st.columns(3)
            if c1.button("Continue"):
                progress["paused"] = False
                progress["next_round_note"] = note.strip()
                st.session_state.in_progress = progress
                st.rerun()
            if c2.button("Stop here"):
                progress["rounds"] = completed
                progress["paused"] = False
                st.session_state.in_progress = progress
                st.rerun()
            if c3.button("Cancel this question", key="debate_cancel"):
                st.session_state.in_progress = None
                st.rerun()
            st.stop()

        while completed < rounds:
            i = completed
            speaker_name = "Claude" if speaker == "claude" else "ChatGPT"
            role_instruction = all_roles.get(roles[speaker], "Continue the discussion in your assigned role.")
            if progress.get("steelman") and i > 0:
                role_instruction = f"{STEELMAN_INSTRUCTION}\n\n{role_instruction}"
            if progress.get("next_round_note"):
                role_instruction += f"\n\nAdditional note from the user: {progress['next_round_note']}"
                progress["next_round_note"] = ""
            message = build_debate_message(prompt, transcript, role_instruction)
            st.caption(f"Round {i + 1} of {rounds} - {speaker_name} as {roles[speaker]}")  # Step: stage indicator
            reply = run_step(
                f"round_{i}", f"**{speaker_name}** (round {i + 1}, {roles[speaker]}):",
                (lambda m=message, s=speaker: (ask_claude_stream if s == "claude" else ask_chatgpt_stream)(
                    [{"role": "user", "content": m}], system=sys_prompt, max_tokens=mt)),
            )
            transcript.append((speaker_name, reply))
            speaker = "chatgpt" if speaker == "claude" else "claude"
            completed += 1

            # Adaptive debate length: a cheap yes/no check, only once there's
            # enough transcript to judge, and only if there's still a round
            # left to potentially skip - never runs past the chosen cap,
            # only ever stops the debate earlier than it.
            if length_mode != "Fixed rounds" and completed >= 2 and completed < rounds:
                if length_mode.startswith("Stop early"):
                    last_two_text = "\n\n".join(f"{s}: {r}" for s, r in transcript[-2:])
                    stalled = run_quiet_step(
                        f"stall_check_{completed}",
                        (lambda t=last_two_text: ask_claude(
                            [{"role": "user", "content": build_stall_check_prompt(t)}], system=sys_prompt, max_tokens=10)),
                    )
                    if "yes" in stalled.lower():
                        st.caption(f"Stopping early after round {completed} - replies had stopped adding new information.")
                        progress["rounds"] = completed
                        rounds = completed
                else:  # consensus-or-bust
                    transcript_text = "\n\n".join(f"{s}: {r}" for s, r in transcript)
                    agreed = run_quiet_step(
                        f"consensus_check_{completed}",
                        (lambda t=transcript_text: ask_claude(
                            [{"role": "user", "content": build_consensus_check_prompt(t)}], system=sys_prompt, max_tokens=10)),
                    )
                    if "yes" in agreed.lower():
                        st.caption(f"Stopping after round {completed} - both sides converged on a shared conclusion.")
                        progress["rounds"] = completed
                        rounds = completed

            if pause_on and completed < rounds:
                progress["paused"] = True
                st.session_state.in_progress = progress
                st.rerun()

        new_turn = {
            "type": "debate", "question": question, "transcript": transcript,
            "claude_role": progress["claude_role"], "chatgpt_role": progress["chatgpt_role"],
        }

    elif mode.startswith("Recipe"):
        transcript = []
        for i, (speaker, label, instruction) in enumerate(RECIPE_STEPS):
            speaker_name = "Claude" if speaker == "claude" else "ChatGPT"
            so_far = [(s, r) for s, _l, r in transcript]
            message = build_debate_message(prompt, so_far, instruction)
            st.caption(f"Step {i + 1} of {len(RECIPE_STEPS)} - {speaker_name} ({label})")  # Step: stage indicator
            reply = run_step(
                f"recipe_{i}", f"**{speaker_name}** ({label}):",
                (lambda m=message, s=speaker: (ask_claude_stream if s == "claude" else ask_chatgpt_stream)(
                    [{"role": "user", "content": m}], system=sys_prompt, max_tokens=mt)),
            )
            transcript.append((speaker_name, label, reply))
        new_turn = {"type": "recipe", "question": question, "transcript": transcript}

    else:  # One model only
        which_model = progress["which_model"]
        history_ctx = chat["history"] + [{"role": "user", "content": prompt}]
        answer = run_step(
            "answer", f"**{which_model}:**",
            lambda: (ask_claude_stream if which_model == "Claude" else ask_chatgpt_stream)(history_ctx, system=sys_prompt, max_tokens=mt),
        )
        new_turn = {"type": "solo", "question": question, "model": which_model, "answer": answer}

    # Every step succeeded - persist the turn and clear the in-progress state.
    # Always the plain-text version here, never the raw `prompt` - that may be
    # a list of content blocks (text + image) when an image was attached, and
    # only the text belongs in permanent, resend-on-every-follow-up context.
    new_turn["confidence_tags"] = progress.get("confidence_tags", False)
    new_turn["blind_judging"] = progress.get("blind_judging", False)

    # "Edit my file(s)": every mode now has a way to get one or more
    # candidate edited files back out of it.
    #  - solo / single review: one source (the reviewer's answer, for single
    #    review, since that's the one built on top of a critique).
    #  - independent answers: TWO sources - Claude's and ChatGPT's answers
    #    were both already asked for the edited file(s), independently.
    #  - debate / recipe: neither has one clean "final answer" mid-flow, and
    #    asking for the whole file back on every round/step would be a
    #    wasteful, repeated cost - so one extra quiet call runs here instead,
    #    using the whole discussion as context, only when file_edit_mode is on.
    new_turn["file_edits"] = []
    if progress.get("file_edit_mode") and progress.get("file_edit_filenames"):
        filenames = progress["file_edit_filenames"]

        if new_turn["type"] == "solo":
            for fname, edited_text in extract_edited_files(new_turn["answer"], filenames).items():
                new_turn["file_edits"].append({"source": new_turn["model"], "filename": fname, "edited_text": edited_text})

        elif new_turn["type"] == "single":
            reviewer_key = "chatgpt" if new_turn.get("primary", "Claude") == "Claude" else "claude"
            reviewer_name = "ChatGPT" if reviewer_key == "chatgpt" else "Claude"
            for fname, edited_text in extract_edited_files(new_turn[reviewer_key], filenames).items():
                new_turn["file_edits"].append({"source": reviewer_name, "filename": fname, "edited_text": edited_text})

        elif new_turn["type"] == "independent":
            for model_key, model_name in (("claude", "Claude"), ("chatgpt", "ChatGPT")):
                for fname, edited_text in extract_edited_files(new_turn[model_key], filenames).items():
                    new_turn["file_edits"].append({"source": model_name, "filename": fname, "edited_text": edited_text})

        elif new_turn["type"] in ("debate", "recipe"):
            discussion = turn_transcript_text(new_turn)
            fe_prompt = f"Question: {question}\n\nDiscussion so far:\n{discussion}\n\n{build_file_edit_instruction(filenames)}"
            st.caption("Preparing the final edited file(s) from the discussion above...")
            fe_answer = run_quiet_step(
                "file_edit_final",
                lambda: ask_claude([{"role": "user", "content": fe_prompt}], system=memory),
            )
            for fname, edited_text in extract_edited_files(fe_answer, filenames).items():
                new_turn["file_edits"].append({"source": None, "filename": fname, "edited_text": edited_text})

    db.add_turn(chat_id, flow_turn_index, new_turn, progress.get("prompt_for_storage", prompt))
    for edit in new_turn["file_edits"]:
        db.add_extra(chat_id, flow_turn_index, new_turn["type"], "file_edit", edit["source"], edit["edited_text"],
                      in_context=False, role=edit["filename"])
    if chat["title"] == "New chat":
        db.rename_chat(chat_id, question[:40])
    st.session_state.in_progress = None
    st.rerun()

st.divider()

# Show the conversation so far, oldest first, each answer with its own
# download button. Older chats saved before this mode toggle existed have
# no "type" key, so they're treated as "single" too.
last_index = len(chat["display"]) - 1

for i, turn in enumerate(chat["display"]):
    st.markdown(f"**You:** {turn['question']}")

    turn_type = turn.get("type", "single")
    tagged = turn.get("confidence_tags", False)

    # Quick win: copy as Markdown / LaTeX - one small control per turn that
    # switches every answer block below between the normal rendered view
    # and a copy-friendly st.code() block (which gets a copy icon for free).
    view_mode = st.radio("View as", ["Rendered", "Markdown", "LaTeX"], horizontal=True,
                         key=f"view_{i}", label_visibility="collapsed")

    if turn_type == "single":
        primary = turn.get("primary", "Claude")  # chats saved before this existed default to Claude-first
        claude_label = "answered first" if primary == "Claude" else "reviewed"
        chatgpt_label = "answered first" if primary == "ChatGPT" else "reviewed"

        st.markdown(f"**Claude** ({claude_label}):")
        show_text(turn["claude"], view_mode, tagged)
        st.download_button("Download Claude's answer", turn["claude"], file_name=f"claude_answer_{i}.txt", key=f"dl_claude_{i}")

        st.markdown(f"**ChatGPT** ({chatgpt_label}):")
        show_text(turn["chatgpt"], view_mode, tagged)
        st.download_button("Download ChatGPT's answer", turn["chatgpt"], file_name=f"chatgpt_answer_{i}.txt", key=f"dl_chatgpt_{i}")

    elif turn_type == "independent":
        # Quick win: blind judging - only for turns asked with it on, and
        # only until a vote is cast (or skipped). Reusing a per-turn
        # deterministic shuffle (seeded from chat_id+turn index) means the
        # A/B assignment stays the same across reruns without storing it.
        existing_vote = db.get_vote(chat_id, i) if turn.get("blind_judging") else "revealed"

        if existing_vote is None:
            order = ["Claude", "ChatGPT"]
            random.Random(f"{chat_id}:{i}").shuffle(order)
            answer_by_model = {"Claude": turn["claude"], "ChatGPT": turn["chatgpt"]}
            col_a, col_b = st.columns(2)
            with col_a:
                st.markdown("**Answer A:**")
                show_text(answer_by_model[order[0]], view_mode, tagged)
            with col_b:
                st.markdown("**Answer B:**")
                show_text(answer_by_model[order[1]], view_mode, tagged)

            st.caption("Blind judging - vote before the model names are revealed.")
            v1, v2, v3, v4 = st.columns(4)
            if v1.button("Answer A is better", key=f"vote_a_{i}"):
                db.add_vote(chat_id, i, order[0]); st.rerun()
            if v2.button("Answer B is better", key=f"vote_b_{i}"):
                db.add_vote(chat_id, i, order[1]); st.rerun()
            if v3.button("Tie", key=f"vote_tie_{i}"):
                db.add_vote(chat_id, i, "Tie"); st.rerun()
            if v4.button("Reveal without voting", key=f"vote_skip_{i}"):
                db.add_vote(chat_id, i, "(skipped)"); st.rerun()

        else:
            # Quick win 4: the two independent answers side by side, so the
            # differences the comparison talks about are easy to scan across.
            if existing_vote not in (None, "revealed"):
                st.caption(f"Blind vote: {existing_vote}" if existing_vote != "(skipped)" else "Blind judging skipped.")
            col_a, col_b = st.columns(2)
            with col_a:
                st.markdown("**Claude** (answered independently):")
                show_text(turn["claude"], view_mode, tagged)
                st.download_button("Download Claude's answer", turn["claude"], file_name=f"claude_answer_{i}.txt", key=f"dl_ind_claude_{i}")
            with col_b:
                st.markdown("**ChatGPT** (answered independently):")
                show_text(turn["chatgpt"], view_mode, tagged)
                st.download_button("Download ChatGPT's answer", turn["chatgpt"], file_name=f"chatgpt_answer_{i}.txt", key=f"dl_ind_chatgpt_{i}")

            st.markdown("**Key differences & comparison:**")
            show_text(turn["comparison"], view_mode, tagged)
            st.download_button("Download comparison", turn["comparison"], file_name=f"comparison_{i}.txt", key=f"dl_ind_compare_{i}")

    elif turn_type == "solo":
        st.markdown(f"**{turn['model']}:**")
        show_text(turn["answer"], view_mode, tagged)
        st.download_button(f"Download {turn['model']}'s answer", turn["answer"], file_name=f"{turn['model'].lower()}_answer_{i}.txt", key=f"dl_solo_{i}")

    else:  # debate or recipe
        for j, entry in enumerate(turn["transcript"]):
            if turn_type == "recipe":
                speaker, label, reply = entry
                header = f"**{speaker} ({label}):**"
                dl_label = f"Download step {j + 1} ({speaker})"
            else:
                speaker, reply = entry
                role = turn.get("claude_role") if speaker == "Claude" else turn.get("chatgpt_role")
                role_label = f", {role}" if role else ""
                header = f"**{speaker} (round {j + 1}{role_label}):**"
                dl_label = f"Download round {j + 1} ({speaker})"
            st.markdown(header)
            show_text(reply, view_mode, tagged)
            st.download_button(dl_label, reply, file_name=f"{turn_type}_{i}_{j + 1}_{speaker.lower()}.txt", key=f"dl_{turn_type}_{i}_{j}")

    if turn.get("conclusion"):
        st.markdown("**Conclusion:**")
        show_text(turn["conclusion"], view_mode, tagged)
        st.download_button("Download conclusion", turn["conclusion"], file_name=f"conclusion_{i}.txt", key=f"dl_conclusion_{i}")

    if turn.get("disagreement_map"):
        st.markdown("**Disagreement map:**")
        show_text(turn["disagreement_map"], view_mode, tagged)
        st.download_button("Download disagreement map", turn["disagreement_map"], file_name=f"disagreement_map_{i}.txt", key=f"dl_dmap_{i}")

    if turn.get("third_option"):
        st.markdown("**A third option:**")
        show_text(turn["third_option"], view_mode, tagged)
        st.download_button("Download third option", turn["third_option"], file_name=f"third_option_{i}.txt", key=f"dl_third_{i}")

    for edit_idx, edit in enumerate(turn.get("file_edits", [])):
        original_bytes = db.get_original_file_bytes(chat_id, i, edit["filename"])
        edited_bytes, edited_dl_name, edited_mime = build_edited_file_bytes(
            edit["edited_text"], edit["filename"], original_bytes=original_bytes,
        )
        source_label = f" ({edit['source']}'s version)" if edit.get("source") else ""
        st.download_button(
            f'Download edited "{edited_dl_name}"{source_label}', edited_bytes,
            file_name=edited_dl_name, mime=edited_mime, key=f"dl_editedfile_{i}_{edit_idx}",
        )

    # Step 22/30 + smarter-collaboration batch: available on every past turn -
    # none of these depend on what happened afterward. Disagreement map and
    # "suggest a third option" only make sense where two models actually had
    # a discussion, so they're left off solo turns.
    can_act = st.session_state.in_progress is None
    action_specs = [("Get final conclusion", "conclusion", bool(turn.get("conclusion")))]
    if turn_type != "solo":
        action_specs.append(("Disagreement map", "disagreement_map", bool(turn.get("disagreement_map"))))
        action_specs.append(("Suggest a third option", "third_option", bool(turn.get("third_option"))))
    action_specs.append(("Add to memory", "memory", False))
    action_specs.append(("Branch from here", "branch", False))

    action_cols = st.columns(len(action_specs))
    for col, (label, kind, already_done) in zip(action_cols, action_specs):
        if col.button(label, key=f"{kind}_{i}", disabled=not can_act or already_done):
            if kind == "branch":
                new_chat_id = db.branch_chat(chat_id, i)
                st.session_state.current_chat_id = new_chat_id
                st.session_state.in_progress = None
            else:
                st.session_state.pending_action = {"kind": kind, "turn_index": i}
            st.rerun()

    # Decision journal: log what you chose and why straight from the turn
    # that prompted it; the outcome (what actually happened) gets filled in
    # later, here or in the sidebar's Decision journal panel.
    with st.expander("Log a decision from this turn"):
        logged = [d for d in chat_decisions if d["turn_index"] == i]
        for d in logged:
            outcome_note = f" (outcome: {d['outcome']})" if d.get("outcome") else ""
            st.caption(f"Logged: **{d['choice']}**" + (f" - {d['reasoning']}" if d["reasoning"] else "") + outcome_note)
        new_choice = st.text_input("What did you decide?", key=f"decision_choice_{i}")
        new_reasoning = st.text_area("Why (optional)", key=f"decision_reasoning_{i}")
        if st.button("Save decision", key=f"decision_save_{i}") and new_choice.strip():
            db.add_decision(chat_id, i, new_choice.strip(), new_reasoning.strip())
            st.rerun()

    # Step 27: "continue from this answer" - only offered on the most recent
    # turn, and only for the modes that actually feed into the shared
    # history (debate/recipe are their own separate side conversation, so
    # there's nothing coherent for a follow-up to build on there).
    if turn_type in ("solo", "single", "independent") and i == last_index:
        st.caption("Continue from this answer:")
        fc = st.columns(4)
        for col, (action_key, label) in zip(fc, [
            ("explain", "Explain this"),
            ("challenge", "Challenge this"),
            ("practical", "Make it practical"),
            ("other_model", "Ask the other model"),
        ]):
            if col.button(label, key=f"followup_{i}_{action_key}", disabled=not can_act):
                st.session_state.pending_action = {"kind": "followup", "action": action_key, "based_on_index": i, "turn_index": i}
                st.rerun()

    st.divider()
